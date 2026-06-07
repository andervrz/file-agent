# agent/tools/office_tools.py
from __future__ import annotations

import asyncio
import re
from pathlib import Path

from .base import BaseTool, PathSafeguard, ToolResult


# ─── Word ─────────────────────────────────────────────────────────────────────

class CreateWordTool(BaseTool):
    name = "create_word"
    description = "Crea un documento Word (.docx) con contenido formateado."
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "Ruta de salida .docx"},
            "content": {"type": "string", "description": "Contenido en markdown simple"},
            "title": {"type": "string", "description": "Título del documento (opcional)"},
        },
        "required": ["path", "content"],
    }

    def __init__(self, safeguard: PathSafeguard):
        self._safeguard = safeguard

    async def execute(self, tool_use_id: str, **kwargs) -> ToolResult:
        try:
            path = self._safeguard.validate(kwargs["path"])
            content = kwargs.get("content", "")
            title = kwargs.get("title", "")

            if not path.suffix:
                path = path.with_suffix(".docx")

            path.parent.mkdir(parents=True, exist_ok=True)
            await asyncio.to_thread(_create_docx, path, content, title)
            return ToolResult(
                tool_use_id=tool_use_id,
                content=f"Documento Word creado: {path}",
            )
        except ImportError:
            return ToolResult(
                tool_use_id=tool_use_id,
                content="python-docx no está instalado. Ejecuta: uv add python-docx",
                is_error=True,
            )
        except PermissionError as e:
            return ToolResult(tool_use_id=tool_use_id, content=str(e), is_error=True)
        except Exception as e:
            return ToolResult(
                tool_use_id=tool_use_id,
                content=f"Error al crear Word: {e}",
                is_error=True,
            )


def _create_docx(path: Path, content: str, title: str) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    if title:
        doc.add_heading(title, level=0)

    for line in content.splitlines():
        stripped = line.rstrip()

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith(("- ", "* ")):
            p = doc.add_paragraph(stripped[2:], style="List Bullet")
        elif re.match(r"^\d+\. ", stripped):
            text = re.sub(r"^\d+\. ", "", stripped)
            doc.add_paragraph(text, style="List Number")
        elif stripped == "---":
            doc.add_paragraph("─" * 40)
        elif stripped:
            p = doc.add_paragraph()
            # Inline bold: **texto**
            parts = re.split(r"\*\*(.+?)\*\*", stripped)
            for i, part in enumerate(parts):
                run = p.add_run(part)
                run.bold = (i % 2 == 1)
        else:
            doc.add_paragraph("")

    doc.save(str(path))


# ─── Excel ────────────────────────────────────────────────────────────────────

class CreateExcelTool(BaseTool):
    name = "create_excel"
    description = "Crea una hoja de cálculo Excel (.xlsx) con datos en formato CSV."
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "data": {
                "type": "string",
                "description": "Filas CSV separadas por \\n, columnas por coma",
            },
            "sheet_name": {"type": "string", "default": "Sheet1"},
            "headers": {"type": "boolean", "default": True},
        },
        "required": ["path", "data"],
    }

    def __init__(self, safeguard: PathSafeguard):
        self._safeguard = safeguard

    async def execute(self, tool_use_id: str, **kwargs) -> ToolResult:
        try:
            path = self._safeguard.validate(kwargs["path"])
            data = kwargs.get("data", "")
            sheet_name = kwargs.get("sheet_name", "Sheet1")
            headers = kwargs.get("headers", True)

            if not path.suffix:
                path = path.with_suffix(".xlsx")

            path.parent.mkdir(parents=True, exist_ok=True)
            rows, cols = await asyncio.to_thread(
                _create_xlsx, path, data, sheet_name, headers
            )
            return ToolResult(
                tool_use_id=tool_use_id,
                content=f"Hoja Excel creada: {path} ({rows} filas, {cols} columnas)",
            )
        except ImportError:
            return ToolResult(
                tool_use_id=tool_use_id,
                content="openpyxl no está instalado. Ejecuta: uv add openpyxl",
                is_error=True,
            )
        except PermissionError as e:
            return ToolResult(tool_use_id=tool_use_id, content=str(e), is_error=True)
        except Exception as e:
            return ToolResult(
                tool_use_id=tool_use_id,
                content=f"Error al crear Excel: {e}",
                is_error=True,
            )


def _create_xlsx(
    path: Path, data: str, sheet_name: str, headers: bool
) -> tuple[int, int]:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    rows_data = [
        [cell.strip() for cell in line.split(",")]
        for line in data.strip().splitlines()
        if line.strip()
    ]

    if not rows_data:
        wb.save(str(path))
        return 0, 0

    max_cols = max(len(r) for r in rows_data)

    for row_idx, row in enumerate(rows_data, start=1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if row_idx == 1 and headers:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D3D3D3")

    # Autoajustar ancho de columnas
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    wb.save(str(path))
    return len(rows_data), max_cols


# ─── PowerPoint ───────────────────────────────────────────────────────────────

class CreatePresentationTool(BaseTool):
    name = "create_presentation"
    description = "Crea una presentación PowerPoint (.pptx)."
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "slides": {
                "type": "string",
                "description": 'Slides en formato "TITLE: título\\nCONTENT: contenido\\n---\\n..."',
            },
            "theme": {
                "type": "string",
                "enum": ["light", "dark", "minimal"],
                "default": "light",
            },
        },
        "required": ["path", "slides"],
    }

    def __init__(self, safeguard: PathSafeguard):
        self._safeguard = safeguard

    async def execute(self, tool_use_id: str, **kwargs) -> ToolResult:
        try:
            path = self._safeguard.validate(kwargs["path"])
            slides_str = kwargs.get("slides", "")
            theme = kwargs.get("theme", "light")

            if not path.suffix:
                path = path.with_suffix(".pptx")

            path.parent.mkdir(parents=True, exist_ok=True)
            slides = _parse_slides(slides_str)

            if not slides:
                return ToolResult(
                    tool_use_id=tool_use_id,
                    content='Formato incorrecto. Usa "TITLE: título\\nCONTENT: contenido\\n---"',
                    is_error=True,
                )

            n = await asyncio.to_thread(_create_pptx, path, slides, theme)
            return ToolResult(
                tool_use_id=tool_use_id,
                content=f"Presentación creada: {path} ({n} slides)",
            )
        except ImportError:
            return ToolResult(
                tool_use_id=tool_use_id,
                content="python-pptx no está instalado. Ejecuta: uv add python-pptx",
                is_error=True,
            )
        except PermissionError as e:
            return ToolResult(tool_use_id=tool_use_id, content=str(e), is_error=True)
        except Exception as e:
            return ToolResult(
                tool_use_id=tool_use_id,
                content=f"Error al crear presentación: {e}",
                is_error=True,
            )


def _parse_slides(text: str) -> list[dict[str, str]]:
    slides: list[dict[str, str]] = []
    for block in text.split("---"):
        block = block.strip()
        if not block:
            continue
        slide: dict[str, str] = {"title": "", "content": ""}
        for line in block.splitlines():
            if line.upper().startswith("TITLE:"):
                slide["title"] = line.split(":", 1)[1].strip()
            elif line.upper().startswith("CONTENT:"):
                slide["content"] = line.split(":", 1)[1].strip()
        if slide["title"] or slide["content"]:
            slides.append(slide)
    return slides


def _create_pptx(
    path: Path, slides: list[dict[str, str]], theme: str
) -> int:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()

    # Colores por tema
    themes = {
        "light": {"bg": RGBColor(0xFF, 0xFF, 0xFF), "title": RGBColor(0x1A, 0x1A, 0x2E), "body": RGBColor(0x33, 0x33, 0x33)},
        "dark":  {"bg": RGBColor(0x1A, 0x1A, 0x2E), "title": RGBColor(0xFF, 0xFF, 0xFF), "body": RGBColor(0xCC, 0xCC, 0xCC)},
        "minimal": {"bg": RGBColor(0xF8, 0xF8, 0xF8), "title": RGBColor(0x00, 0x00, 0x00), "body": RGBColor(0x44, 0x44, 0x44)},
    }
    colors = themes.get(theme, themes["light"])

    slide_layout = prs.slide_layouts[1]  # título + contenido

    for slide_data in slides:
        slide = prs.slides.add_slide(slide_layout)
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = colors["bg"]

        if slide.shapes.title:
            slide.shapes.title.text = slide_data["title"]
            slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = colors["title"]
            slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
            slide.shapes.title.text_frame.paragraphs[0].runs[0].font.bold = True

        if len(slide.placeholders) > 1:
            body = slide.placeholders[1]
            tf = body.text_frame
            tf.text = slide_data["content"]
            for para in tf.paragraphs:
                for run in para.runs:
                    run.font.color.rgb = colors["body"]
                    run.font.size = Pt(18)

    prs.save(str(path))
    return len(slides)
